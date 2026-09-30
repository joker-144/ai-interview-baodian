"""引擎 A 输入侧：简历上传 → 解析 → 体检 → AI 一键优化 → 多版本管理。

真实链路（一期）：multipart 上传 → 本地抽取纯文本（PDF/DOCX/TXT，见 app/resume_parser.py）
→ 主模型结构化与能力维度评估（app/analyzer.py，批 3 起同次产出体检 checkup）
→ 保存为一个简历版本（store.add_resume_version，内存 + resumes v2 表尽力而为双写）。
同时留存原文与画像摘要供出题提示词使用（不对外回显原文）。

扫描件/图片型 PDF 本地抽不出文字时，自动转 vision 分层多模态 OCR 兜底
（app/resume_ocr.py）；OCR 不可用（分层停用 / 模型不支持）时返回可读提示。

批 3 新增：
- GET /api/resumes 多版本列表（各版本挂体检报告）；
- GET /{resume_id}/checkup 体检报告（旧记录缺 checkup 时惰性补算并写回）；
- POST /{resume_id}/optimize AI 一键优化（主模型重写，生成优化版新版本）；
- GET /{resume_id}/download 导出 DOCX（python-docx，文件名 UTF-8 转义）；
- DELETE /{resume_id} 删除版本。

出题任务不在本路由：统一走 `POST /api/question-sets/generate`（见 question_sets.py），
三通道（简历 / 岗位检索 / JD 定向）共用同一入口，由 body.source 区分；
出题默认用最新一份（/latest 语义保持）。
"""

import io
import urllib.parse

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from app import analyzer, llm, resume_ocr, resume_parser, store
from app.auth import get_current_user
from app.schemas import ResumeAnalysis

router = APIRouter(prefix="/api/resumes", tags=["resumes"])

MAX_FILE_MB = 10
# 单份简历导出最多写入的行数（防御异常超长文本拖垮 DOCX 生成）
DOCX_MAX_LINES = 2000


def _require_version(user_id: str, resume_id: str) -> dict:
    rec = store.get_resume_version(user_id, resume_id)
    if not rec:
        raise HTTPException(status_code=404, detail="简历不存在或已被删除")
    return rec


@router.post("", response_model=ResumeAnalysis)
async def upload_resume(
    file: UploadFile = File(...), user_id: str = Depends(get_current_user)
) -> dict:
    """上传简历并解析（真实解析 + 主模型结构化 + 体检，耗时约 10~30s）。

    扫描件/图片型 PDF 自动转多模态 OCR 再解析，耗时会更长（页图渲染 + 视觉模型转写）。
    """
    file_name = file.filename or "resume"
    data = await file.read()
    if not data:
        raise HTTPException(status_code=422, detail="上传文件为空")
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_FILE_MB}MB 上限")

    try:
        text = resume_parser.extract_text(file_name, data)
    except resume_parser.ScannedDocument as exc:
        # 扫描件/图片型 PDF：转 vision 分层 OCR；不可用时给可读降级提示
        try:
            text = await resume_ocr.ocr_pdf(data)
        except resume_ocr.OcrError as inner:
            raise HTTPException(status_code=422, detail=str(inner)) from inner
        except llm.LlmError as inner:
            raise HTTPException(
                status_code=502,
                detail=(
                    f"扫描件 OCR 失败：{inner}；"
                    "请在管理端检查「多模态模型」分层配置，或上传可复制文字的版本"
                ),
            ) from inner
    except resume_parser.UnsupportedFormat as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    try:
        analysis, summary = await analyzer.analyze_resume(text, file_name)
    except llm.LlmError as exc:
        raise HTTPException(status_code=502, detail=f"简历解析失败：{exc}") from exc

    rec = store.add_resume_version(
        user_id, file_name=file_name, text=text, summary=summary, analysis=analysis
    )
    return {**analysis, "resumeId": rec["resumeId"]}


@router.get("")
def list_resumes(user_id: str = Depends(get_current_user)) -> list[dict]:
    """多版本简历列表（version 降序；analysis 内含体检报告，可能为 None）。"""
    return [
        {
            "resumeId": r["resumeId"],
            "version": r["version"],
            "isOptimized": r["isOptimized"],
            "fileName": r["fileName"],
            "createdAt": r["createdAt"],
            "analysis": r["analysis"],
        }
        for r in store.user_resume_versions(user_id)
    ]


@router.get("/latest", response_model=ResumeAnalysis | None)
def get_latest_analysis(user_id: str = Depends(get_current_user)) -> dict | None:
    """读取最新一份简历解析结果（含能力维度评估与预估题量；出题默认用这份）。"""
    return (store.user_resume(user_id) or {}).get("analysis")


@router.get("/{resume_id}", response_model=ResumeAnalysis)
def get_analysis(resume_id: str, user_id: str = Depends(get_current_user)) -> dict:
    """按 id 读取解析结果。"""
    rec = _require_version(user_id, resume_id)
    if not rec.get("analysis"):
        raise HTTPException(status_code=404, detail="该简历缺少解析结果")
    return {**rec["analysis"], "resumeId": resume_id}


@router.delete("/{resume_id}")
def delete_resume(resume_id: str, user_id: str = Depends(get_current_user)) -> dict:
    """删除一个简历版本（最新版被删时自动回退到次新版本作为出题底稿）。"""
    if not store.drop_resume_version(user_id, resume_id):
        raise HTTPException(status_code=404, detail="简历不存在或已被删除")
    return {"ok": True}


@router.get("/{resume_id}/checkup")
async def get_checkup(resume_id: str, user_id: str = Depends(get_current_user)) -> dict:
    """简历体检报告；旧记录缺 checkup 时惰性补算并写回（LLM 失败走规则兜底）。"""
    rec = _require_version(user_id, resume_id)
    analysis = dict(rec.get("analysis") or {})
    if analysis.get("checkup"):
        return {"checkup": analysis["checkup"]}
    analysis["checkup"] = await analyzer.run_checkup(
        rec.get("text", ""), analysis.get("dimensions") or []
    )
    store.update_resume_version(user_id, resume_id, analysis)
    return {"checkup": analysis["checkup"]}


@router.post("/{resume_id}/optimize", response_model=ResumeAnalysis)
async def optimize_resume(
    resume_id: str, user_id: str = Depends(get_current_user)
) -> dict:
    """AI 一键优化：按体检结论主模型重写，生成优化版新版本（version+1、is_optimized=1）。

    输入 = 体检结论 + 原简历结构化数据（文档 3.4 实现要点）；优化失败返回 502，不静默降级。
    """
    rec = _require_version(user_id, resume_id)
    base = dict(rec.get("analysis") or {})
    if not base.get("dimensions"):
        raise HTTPException(status_code=409, detail="该简历缺少解析结果，无法优化")
    if not base.get("checkup"):
        base["checkup"] = await analyzer.run_checkup(
            rec.get("text", ""), base.get("dimensions") or []
        )
        store.update_resume_version(user_id, resume_id, base)

    try:
        result = await analyzer.optimize_resume(
            rec.get("text", ""), rec.get("summary", ""), base["checkup"]
        )
    except llm.LlmError as exc:
        raise HTTPException(status_code=502, detail=f"AI 优化失败：{exc}") from exc

    optimized_analysis = {
        **base,
        "fileName": f"优化版-{base.get('fileName', '简历')}",
        "checkup": result["checkup"],
    }
    new_rec = store.add_resume_version(
        user_id,
        file_name=optimized_analysis["fileName"],
        text=result["text"],
        summary=rec.get("summary", ""),
        analysis=optimized_analysis,
        is_optimized=True,
    )
    return {**optimized_analysis, "resumeId": new_rec["resumeId"]}


@router.get("/{resume_id}/download")
def download_resume(resume_id: str, user_id: str = Depends(get_current_user)):
    """导出简历 DOCX（python-docx 生成；文件名 UTF-8 转义，兼容中文文件名）。"""
    rec = _require_version(user_id, resume_id)
    text = (rec.get("text") or "").strip()
    if not text:
        raise HTTPException(status_code=404, detail="该简历没有可下载的原文")
    buffer = _build_docx(text)

    raw_name = (rec.get("fileName") or "简历").rsplit(".", 1)[0]
    file_name = f"{raw_name or '简历'}.docx"
    quoted = urllib.parse.quote(file_name)
    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"{quoted}\"; filename*=UTF-8''{quoted}"
            )
        },
    )


def _build_docx(text: str) -> io.BytesIO:
    """纯文本简历 → DOCX：短行识别为标题、·/- 开头识别为列表项，其余为段落。"""
    try:
        from docx import Document  # 延迟导入：未装 python-docx 时其余接口不受影响
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(
            status_code=502, detail="服务端未安装 python-docx，无法导出 DOCX"
        ) from exc

    document = Document()
    for line in text.splitlines()[:DOCX_MAX_LINES]:
        content = line.strip()
        if not content:
            continue
        if content[0] in "·-•*":
            body = content.lstrip("·-•* ").strip()
            if body:
                document.add_paragraph(body, style="List Bullet")
        elif len(content) <= 30 and not content.endswith(("。", "，", "、", "；", ";", ",")):
            document.add_heading(content, level=2)
        else:
            document.add_paragraph(content)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
