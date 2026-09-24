"""简历文件文本抽取（一期支持 PDF / DOCX / TXT，均为纯本地解析、零 API 成本）。

扫描件（图片型 PDF）抽不出文字时，抛 UnsupportedFormat 交给路由层返回可读提示，
多模态 OCR 属后续版本（vision 分层已在管理端预留）。
"""

import io
import re

# PDF 逐字抽取常在中文字符间插入空格（"廖 上 琳"），统一压缩
_CJK = r"\u4e00-\u9fff\u3000-\u303f\uff00-\uffef"
_CJK_GAP_RE = re.compile(rf"(?<=[{_CJK}])\s+(?=[{_CJK}])")
_BLANK_RE = re.compile(r"\n{3,}")
_MIN_CHARS = 30  # 少于该长度视为劣质抽取（多为扫描件）


class UnsupportedFormat(ValueError):
    """文件格式不支持，或内容无法抽出可分析的文字。"""


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(blocks)


def _from_text(data: bytes) -> str:
    for encoding in ("utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedFormat("文本编码无法识别，请另存为 UTF-8 后重试")


def clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CJK_GAP_RE.sub("", text)  # 中文字符间空格
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = _BLANK_RE.sub("\n\n", text)
    return text.strip()


def extract_text(file_name: str, data: bytes) -> str:
    """按扩展名抽取纯文本；失败抛出 UnsupportedFormat（消息可直接展示给用户）。"""
    if not data:
        raise UnsupportedFormat("文件内容为空")

    lower = file_name.lower()
    try:
        if lower.endswith(".pdf"):
            raw = _from_pdf(data)
        elif lower.endswith((".docx", ".doc")):
            raw = _from_docx(data)
        elif lower.endswith((".txt", ".md", ".text")):
            raw = _from_text(data)
        else:
            raise UnsupportedFormat("仅支持 PDF / DOCX / TXT 格式的简历")
    except UnsupportedFormat:
        raise
    except Exception as exc:  # 解析库异常统一转为可读提示
        raise UnsupportedFormat(f"文件解析失败：{exc}") from exc

    text = clean(raw)
    if len(text) < _MIN_CHARS:
        raise UnsupportedFormat(
            "未从文件中抽取到有效文字，可能是扫描件/图片型 PDF；请上传可复制文字的版本"
        )
    return text
